"""Read public data files from the commons into commons.db.

**What this does.** Three loaders, one per dataset, each taking a file that is
already in the commons (brought in by scripts/commons_fetch.py) and writing
its rows into commons.db (commonsdb.py):

  load_pdp        a USDA Pesticide Data Program year zip. Reads the samples
                  and results files LINE BY LINE straight out of the zip —
                  the results file is ~130 MB and 2.8 million rows, and this
                  machine has run out of memory before — keeping only the
                  commodities asked for. Also reads the zip's reference
                  workbook: pesticide names, EPA tolerances, commodity names.
  load_benchmarks EPA's Human Health Benchmarks for Pesticides web page, one
                  row per pesticide with its chronic and acute safe doses.
  load_iris       EPA IRIS's reference-dose table (its Advanced Search page),
                  one row per chemical with its chronic oral reference dose.

Each loader replaces what it loaded before for the same year and commodity,
so running it twice gives the same rows, and returns counts for the ledger
(exposurestore.record_pull). LOADER_VERSIONS says which version of each
loader wrote a pull; bump one when its parsing changes and the next pull
re-reads.

Touches: `commonsdb.py` (where rows go), `commons.py` (checksums),
`scripts/reference_data.py` (the command that runs these),
`tests/test_reference_loaders.py`. Design: docs/exposure.md.

Prompt that produced this file: "pull primary data for 2–3 foods first (kale,
strawberries, potatoes) … raw rows, which dataset/year … lives in exo.db
through the existing stores, not loose files" — adapted so the raw rows live
in commons.db beside the commons files — and "read the zip line by line (the
machine OOM-crashed recently); run one job at a time."
"""
from html.parser import HTMLParser
import io
import json
import re
import zipfile

import commons

LOADER_VERSIONS = {"usda-pdp": 1, "epa-hhbp": 1, "epa-iris-rfd": 1}

# PDP's own column order (its data dictionary); the files have no header row.
SAMPLE_COLUMNS = ("sample_pk", "state", "year_2digit", "month", "day", "site", "commod",
                  "source_id", "variety", "origin", "country", "disttype", "commtype", "claim",
                  "quantity", "growst", "packst", "distst")
RESULT_COLUMNS = ("sample_pk", "commod", "commtype", "lab", "pestcode", "testclass", "concen",
                  "lod", "conunit", "confmethod", "confmethod2", "annotate", "quantitate",
                  "mean", "extract", "determin")

# What to multiply a PDP concentration by to get ppb, by its unit code.
_TO_PPB = {"M": 1000.0, "B": 1.0, "T": 0.001}
_BATCH = 5000


def _member(archive, pattern):
    """The one file in the zip whose name matches `pattern`, or ValueError."""
    matches = [name for name in archive.namelist() if re.search(pattern, name, re.IGNORECASE)]
    if len(matches) != 1:
        raise ValueError(f"expected one file matching {pattern!r} in the zip, found {matches}")
    return matches[0]


def _lines(archive, name):
    """Yield the fields of each line of a pipe-delimited file inside the zip, one at a time."""
    with archive.open(name) as raw:
        for line in io.TextIOWrapper(raw, encoding="latin-1", newline=""):
            line = line.rstrip("\r\n")
            if line:
                yield line.split("|")


def _number(text):
    text = (text or "").strip()
    try:
        return float(text) if text else None
    except ValueError:
        return None


def _integer(text):
    number = _number(text)
    return int(number) if number is not None else None


def _to_ppb(amount, unit):
    if amount is None:
        return None
    return amount * _TO_PPB.get((unit or "").strip().upper(), 1.0)


def _cell_text(value):
    """A workbook cell as text: whole numbers without '.0', empty cells as ''."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _workbook_sheets(archive):
    """The zip's reference workbook as {sheet name: [row as text cells, ...]}.

    Older years ship it as .xls (read with xlrd), 2024 on as .xlsx (openpyxl);
    both are imported only here, so the rest of this file works without them.
    """
    name = _member(archive, r"reference.*\.xlsx?$")
    data = archive.read(name)
    if name.lower().endswith(".xlsx"):
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return {sheet: [[_cell_text(cell) for cell in row]
                        for row in book[sheet].iter_rows(values_only=True)]
                for sheet in book.sheetnames}
    import xlrd
    book = xlrd.open_workbook(file_contents=data)
    return {sheet.name: [[_cell_text(cell) for cell in sheet.row_values(index)]
                         for index in range(sheet.nrows)]
            for sheet in book.sheets()}


def _sheet_rows(sheets, sheet_name, header_first_cell):
    """The rows of one reference sheet after its header row (empty rows skipped)."""
    rows, seen_header = [], False
    for values in sheets.get(sheet_name, []):
        if not seen_header:
            seen_header = bool(values) and values[0].lower() == header_first_cell.lower()
            continue
        if values and values[0]:
            rows.append(values + [""] * 6)
    return rows


def _load_reference(conn, archive, year):
    """Pesticide names, EPA tolerances and commodity names from the zip's workbook.
    Returns (pesticides, tolerances) row counts. Loaded whole: they're small."""
    sheets = _workbook_sheets(archive)
    pesticides = [(year, row[0].zfill(3), row[1], row[2])
                  for row in _sheet_rows(sheets, "Pest Code", "Pest Code")]
    tolerances = [(year, row[0].zfill(3), row[1], row[2], row[3], row[4], row[5])
                  for row in _sheet_rows(sheets, "Tolerance", "Pesticide Code")]
    commodities = [(year, row[0], row[1], _integer(row[2]))
                   for row in _sheet_rows(sheets, "Commodity", "Commodity Code")]
    # Replace this year's reference rows whole, so a re-run never doubles them.
    for table in ("pdp_pesticides", "pdp_tolerances", "pdp_commodities"):
        conn.execute(f"DELETE FROM {table} WHERE year = ?", (year,))
    conn.executemany("INSERT OR REPLACE INTO pdp_pesticides VALUES (?,?,?,?)", pesticides)
    conn.executemany("INSERT OR REPLACE INTO pdp_tolerances VALUES (?,?,?,?,?,?,?)", tolerances)
    conn.executemany("INSERT OR REPLACE INTO pdp_commodities VALUES (?,?,?,?)", commodities)
    return len(pesticides), len(tolerances)


def load_pdp(conn, zip_path, year, commodities):
    """Load one PDP year's samples and results for `commodities` (codes like 'PO').

    Returns {commodity: {"samples", "organic", "conventional", "results"}},
    plus a "_reference" entry with the pesticide and tolerance counts.
    Runs inside the caller's transaction on `conn` (a commons.db connection).
    """
    wanted = {code.upper() for code in commodities}
    counts = {code: {"samples": 0, "organic": 0, "conventional": 0, "results": 0} for code in wanted}
    with zipfile.ZipFile(zip_path) as archive:
        pesticides, tolerances = _load_reference(conn, archive, year)
        # Replace what an earlier pull of these commodities loaded.
        for code in wanted:
            conn.execute("DELETE FROM pdp_samples WHERE year = ? AND commod = ?", (year, code))
            conn.execute("DELETE FROM pdp_results WHERE year = ? AND commod = ?", (year, code))

        # Samples: one small file, kept only for the wanted commodities.
        batch = []
        for fields in _lines(archive, _member(archive, r"samples\.txt$")):
            # Refuse a file laid out differently from the data dictionary this
            # was written against, rather than filing its columns in the wrong places.
            if not 14 <= len(fields) <= len(SAMPLE_COLUMNS) + 1:
                raise ValueError(f"PDP {year} samples line has {len(fields)} fields; expected {len(SAMPLE_COLUMNS)}")
            row = dict(zip(SAMPLE_COLUMNS, fields + [""] * (len(SAMPLE_COLUMNS) - len(fields))))
            code = row["commod"].strip().upper()
            if code not in wanted:
                continue
            claim = row["claim"].strip().upper()
            counts[code]["samples"] += 1
            counts[code]["organic" if claim == "PO" else "conventional"] += 1
            batch.append((year, int(row["sample_pk"]), row["state"], row["month"], row["day"],
                          row["site"], code, row["source_id"].strip(), row["variety"],
                          row["origin"], row["country"], row["disttype"], row["commtype"].strip(),
                          claim, _integer(row["quantity"]), row["growst"], row["packst"],
                          row["distst"]))
        conn.executemany("INSERT OR REPLACE INTO pdp_samples VALUES"
                         " (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", batch)

        # Results: the big file, streamed and inserted in batches so memory stays flat.
        batch = []
        for fields in _lines(archive, _member(archive, r"results\.txt$")):
            if len(fields) < 9:
                continue
            code = fields[1].strip().upper()
            if code not in wanted:
                continue
            row = dict(zip(RESULT_COLUMNS, fields + [""] * (len(RESULT_COLUMNS) - len(fields))))
            concen, lod, unit = _number(row["concen"]), _number(row["lod"]), row["conunit"].strip()
            batch.append((year, int(row["sample_pk"]), code, row["commtype"].strip(), row["lab"],
                          row["pestcode"].strip().zfill(3), row["testclass"], concen, lod, unit,
                          _to_ppb(concen, unit), _to_ppb(lod, unit), row["confmethod"],
                          row["confmethod2"], row["annotate"].strip(), row["quantitate"],
                          row["mean"].strip(), row["extract"], row["determin"]))
            counts[code]["results"] += 1
            if len(batch) >= _BATCH:
                conn.executemany(_RESULT_INSERT, batch)
                batch = []
        if batch:
            conn.executemany(_RESULT_INSERT, batch)
    counts["_reference"] = {"pesticides": pesticides, "tolerances": tolerances}
    return counts


_RESULT_INSERT = "INSERT INTO pdp_results VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)"


class _TableParser(HTMLParser):
    """Collect every table row as a list of (cell text, first link in the cell)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows, self._row, self._cell, self._link = [], None, None, None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell, self._link = [], None
        elif tag == "a" and self._cell is not None and self._link is None:
            self._link = dict(attrs).get("href")

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            self._row.append((" ".join("".join(self._cell).split()), self._link))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None


def _first_number(text):
    match = re.search(r"\d+(?:\.\d+)?(?:[eE]-?\d+)?", (text or "").replace(",", ""))
    return float(match.group(0)) if match else None


def parse_benchmarks(html):
    """EPA's benchmark page → one dict per pesticide (name, cas, doses, memo link).

    Columns are found by their header words, not their position, so a column
    added to the page doesn't shift the numbers into the wrong field.
    """
    parser = _TableParser()
    parser.feed(html)
    header = next((row for row in parser.rows
                   if any("CAS" in cell for cell, _ in row)), None)
    if header is None:
        raise ValueError("no benchmark table with a CAS column on the page")
    labels = [cell.lower() for cell, _ in header]

    def column(*words):
        for index, label in enumerate(labels):
            if all(word in label for word in words):
                return index
        raise ValueError(f"no column with {words} in {labels}")

    name_at, cas_at = 0, column("cas")
    acute_at, chronic_at = column("acute", "pad"), column("chronic", "pad")
    slope_at = column("cancer", "csf")
    out = []
    for row in parser.rows:
        if row is header or len(row) < len(labels):
            continue
        cells = [cell for cell, _ in row]
        if not cells[name_at]:
            continue
        out.append({
            "name": cells[name_at], "cas": cells[cas_at] or None,
            "acute_dose": _first_number(cells[acute_at]),
            "chronic_dose": _first_number(cells[chronic_at]),
            "cancer_slope": _first_number(cells[slope_at]),
            "memo_url": row[name_at][1],
            "as_printed": dict(zip([cell for cell, _ in header], cells)),
        })
    return out


def load_benchmarks(conn, html_path):
    """Load EPA's benchmark page into epa_benchmarks, replacing what was there.
    Returns the number of pesticides loaded."""
    with open(html_path, encoding="utf-8", errors="replace") as handle:
        rows = parse_benchmarks(handle.read())
    checksum = commons.sha256_of(html_path)
    conn.execute("DELETE FROM epa_benchmarks")
    conn.executemany(
        "INSERT OR REPLACE INTO epa_benchmarks (name, cas, acute_dose, chronic_dose, cancer_slope,"
        " memo_url, as_printed, file_sha256) VALUES (?,?,?,?,?,?,?,?)",
        [(row["name"], row["cas"], row["acute_dose"], row["chronic_dose"], row["cancer_slope"],
          row["memo_url"], json.dumps(row["as_printed"], ensure_ascii=False), checksum)
         for row in rows])
    return len(rows)


IRIS_BASE = "https://iris.epa.gov"


def parse_iris_rfd(html):
    """IRIS's reference-dose page → one dict per chemical (name, short name, CAS,
    RfD in mg/kg/day, critical effect, confidence, its IRIS page).

    Columns are found by their header words, as in parse_benchmarks. A row
    whose RfD isn't in mg/kg-day is skipped rather than guessed at.
    """
    parser = _TableParser()
    parser.feed(html)
    header = next((row for row in parser.rows
                   if any(cell.upper() == "RFD VALUE" for cell, _ in row)), None)
    if header is None:
        raise ValueError("no table with an RFD VALUE column on the page")
    labels = [cell.upper() for cell, _ in header]
    name_at, cas_at, rfd_at = labels.index("CHEMICAL NAME"), labels.index("CASRN"), labels.index("RFD VALUE")
    effect_at = labels.index("PRINCIPAL CRITICAL DESCRIPTION")
    confidence_at = labels.index("OVERALL CONFIDENCE")
    out = []
    for row in parser.rows:
        if row is header or len(row) <= max(rfd_at, confidence_at):
            continue
        cells = [cell for cell, _ in row]
        rfd = _first_number(cells[rfd_at])
        if not cells[name_at] or rfd is None or "mg/kg" not in cells[rfd_at]:
            continue
        # "p,p'-Dichlorodiphenyltrichloroethane (DDT)" → short name "DDT".
        short = re.search(r"\(([^()]+)\)\s*$", cells[name_at])
        link = row[name_at][1]
        out.append({
            "name": cells[name_at], "short_name": short.group(1) if short else None,
            "cas": cells[cas_at] or None, "rfd": rfd,
            "critical_effect": cells[effect_at] or None,
            "confidence": cells[confidence_at] or None,
            "landing_url": IRIS_BASE + link if link and link.startswith("/") else link,
            "as_printed": dict(zip([cell for cell, _ in header], cells)),
        })
    return out


def load_iris(conn, html_path):
    """Load IRIS's reference-dose page into iris_rfd, replacing what was there.
    Returns the number of chemicals loaded."""
    with open(html_path, encoding="utf-8", errors="replace") as handle:
        rows = parse_iris_rfd(handle.read())
    checksum = commons.sha256_of(html_path)
    conn.execute("DELETE FROM iris_rfd")
    conn.executemany(
        "INSERT OR REPLACE INTO iris_rfd (name, short_name, cas, rfd, critical_effect, confidence,"
        " landing_url, as_printed, file_sha256) VALUES (?,?,?,?,?,?,?,?,?)",
        [(row["name"], row["short_name"], row["cas"], row["rfd"], row["critical_effect"],
          row["confidence"], row["landing_url"], json.dumps(row["as_printed"], ensure_ascii=False),
          checksum) for row in rows])
    return len(rows)
